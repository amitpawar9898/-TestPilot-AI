"""Build and validate the CA-3 presentation (requires python-pptx 1.0.2)."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation" / "TestPilot_CA3.pptx"
NAVY = RGBColor(15, 23, 42)
INK = RGBColor(30, 41, 59)
MUTED = RGBColor(71, 85, 105)
TEAL = RGBColor(13, 148, 136)
PALE = RGBColor(240, 253, 250)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(248, 250, 252)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, color, radius=True):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    if radius:
        sh.adjustments[0] = 0.12
    return sh


def text(slide, x, y, w, h, value, size=18, color=INK, bold=False, font="Aptos", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.03); tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.01)
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = font; p.font.size = Pt(size); p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(9)
        if align is not None: p.alignment = align
    return box


def base(title, kicker, number):
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
    text(slide, .62, .28, 10.5, .28, kicker.upper(), 10, TEAL, True)
    text(slide, .62, .63, 12, .65, title, 28, NAVY, True)
    rect(slide, .62, 1.4, 12.05, .025, RGBColor(203, 213, 225), False)
    text(slide, 12.2, 7.12, .45, .2, f"{number:02d}", 9, MUTED, True, align=PP_ALIGN.RIGHT)
    return slide


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


# 1. Title
s = prs.slides.add_slide(blank)
s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
rect(s, .8, .8, .13, 4.85, TEAL, False)
text(s, 1.25, 1.05, 10.5, .4, "CA-3 · AGENTIC SYSTEMS", 13, RGBColor(94, 234, 212), True)
text(s, 1.25, 1.85, 10.8, 1.25, "TestPilot AI", 40, WHITE, True)
text(s, 1.27, 3.05, 10.2, 1.05, "Agentic software testing\nand bug detection", 25, RGBColor(203, 213, 225))
rect(s, 1.27, 4.7, 4.35, .58, RGBColor(30, 41, 59))
text(s, 1.47, 4.83, 3.95, .28, "Team: [Add team names]  ·  [Course / date]", 12, WHITE)
text(s, 1.27, 6.8, 8, .25, "Python · FastAPI · React · SQLite", 12, RGBColor(148, 163, 184))
notes(s, "Hello. Our project is TestPilot AI, an agentic software testing and bug-detection platform for Python modules. It coordinates a bounded sequence of specialized agents, records what each did, and reports the evidence it actually collected. The system keeps the original file untouched and can propose a verified patch. I will cover the problem, the agent workflow, the architecture, and measured results from the three bundled samples. Please replace the team and course placeholders before presenting.")

# 2. Problem
s = base("The testing gap", "Problem and domain", 2)
rect(s, .75, 1.8, 5.75, 4.65, WHITE)
text(s, 1.05, 2.08, 5.1, .42, "Why this matters", 19, NAVY, True)
text(s, 1.05, 2.72, 5.05, 2.3, "• Edge cases are easy to miss\n• Test writing takes focused time\n• Generated tests need trustworthy expectations\n• A suggested fix must not silently break passing behavior", 18, INK)
rect(s, 6.85, 1.8, 5.72, 4.65, PALE)
text(s, 7.2, 2.08, 4.9, .42, "Chosen domain: Python QA", 19, NAVY, True)
text(s, 7.2, 2.72, 4.8, 2.3, "Input: source + requirements\nProcess: analyze → test → review\nOutput: evidence + findings + patch diff", 19, INK)
text(s, 7.2, 5.4, 4.75, .7, "Design rule: make test status and oracle quality explicit.", 15, TEAL, True)
notes(s, "The problem is feedback speed and coverage. A developer may know the intended behavior but not have time to explore every boundary or invalid input. A generator can help, but it must distinguish examples with explicit expected values from heuristic robustness checks. TestPilot AI takes Python source and optional requirements, then returns generated cases, observed outcomes, findings, and possibly a patch. The chosen domain fits a multi-agent workflow because planning, analysis, test design, execution, review, repair, and reporting have clear responsibilities and evidence boundaries.")

# 3. Existing and proposed
s = base("From single app to contract-driven system", "Existing challenge → proposal", 3)
rect(s, .78, 1.85, 5.6, 4.7, WHITE)
rect(s, 6.95, 1.85, 5.6, 4.7, PALE)
text(s, 1.12, 2.15, 4.9, .42, "Repository foundation", 18, NAVY, True)
text(s, 1.12, 2.8, 4.8, 2.65, "Streamlit UI\nShared orchestration core\nSQLite persistence\nSandbox-aware runner\nFix proposal and report tools", 17, INK)
text(s, 7.28, 2.15, 4.9, .42, "CA-3 system", 18, NAVY, True)
text(s, 7.28, 2.8, 4.8, 2.65, "React + TypeScript client\nFastAPI JSON contract\n54 checked API operations\nSeven registered agents\nBrowser E2E and evidence docs", 17, INK)
text(s, 5.82, 4.05, .7, .5, "→", 26, TEAL, True, align=PP_ALIGN.CENTER)
notes(s, "The starting repository already had the domain workflow and the seven-agent core in a Streamlit application. The CA-3 work preserves the domain logic while adding a React and TypeScript frontend, a typed FastAPI boundary, and live route and browser verification. The API contract is checked against all 54 method and path operations, and the browser flow exercises registration, sign-in, sample analysis, fix review, downloads, and reports. This is an extension around the existing system rather than a new QA engine invented for the presentation.")

# 4. Agent workflow
s = base("Seven agents, one bounded workflow", "Cooperation and control", 4)
steps = [("1", "Planner"), ("2", "Code Analysis"), ("3", "Test Generation"), ("4", "Test Execution"), ("5", "Critic"), ("6", "Fix"), ("7", "Reporting")]
positions = [(.9, 2.0), (4.65, 2.0), (8.4, 2.0), (8.4, 4.05), (4.65, 4.05), (.9, 4.05), (4.65, 5.35)]
for (n, label), (x, y) in zip(steps, positions):
    rect(s, x, y, 3.05, .9, WHITE)
    text(s, x+.17, y+.14, .42, .4, n, 17, TEAL, True)
    text(s, x+.65, y+.17, 2.2, .4, label, 16, NAVY, True)
text(s, 1.0, 6.55, 11, .4, "Critic may request another evidence cycle only while the configured iteration bound allows it.", 14, MUTED, True)
notes(s, "The sequence is explicit in the orchestrator. Planner checks objective and tool availability. Code Analysis parses the module and produces static findings. Test Generation combines deterministic cases with requirement examples and doctests. Test Execution runs them through the configured runner. The Critic separates confirmed failures from heuristic evidence and can ask for a bounded follow-up cycle. Fix evaluates candidate changes against the same tests, and Reporting persists the merged evidence. Tool calls include names, summaries, results, success state, and duration, so this is inspectable cooperation rather than an opaque agent chat.")

# 5. Architecture
s = base("System architecture", "Components and data flow", 5)
try:
    s.shapes.add_picture(str(ROOT / "docs" / "architecture.svg"), Inches(.85), Inches(1.72), width=Inches(11.65))
except Exception:
    text(s, 1.0, 2.0, 11, 3.8, "Browser → React/Vite → FastAPI → background worker → orchestrator\n\nSeven agents use analysis, test generation, isolated execution, fix verification, and report tools. SQLite stores accounts and run evidence; reports are downloadable through the API.", 21, INK)
notes(s, "This diagram follows the running implementation. The browser uses React and the Vite development proxy to call FastAPI. The API validates and authorizes requests, records run state, then dispatches work to a background worker. The orchestrator shares structured state among the seven agents and invokes domain tools. SQLite stores run and account evidence. Reports and downloads return through typed API routes. Docker is the recommended code-execution boundary; local trusted mode is explicitly not a sandbox. The optional model adapter is separate and requires configuration and consent for remote use.")

# 6. Demo screenshot
s = base("Live flow: sample to verified patch", "Demo evidence", 6)
rect(s, .72, 1.72, 8.12, 4.95, WHITE)
s.shapes.add_picture(str(ROOT / "docs" / "screenshots" / "fix_my_code_sample.png"), Inches(.9), Inches(1.9), width=Inches(7.75))
rect(s, 9.15, 1.82, 3.25, 4.55, PALE)
text(s, 9.48, 2.12, 2.6, .42, "Demo path", 18, NAVY, True)
text(s, 9.48, 2.78, 2.6, 2.7, "Load sample\nRun analysis\nInspect tests\nReview diff\nDownload report", 16, INK)
notes(s, "This screenshot is from the running React application during the browser end-to-end flow. In a live demo, open Fix my code, load the planted-bug sample, and run the analysis. Show the generated tests and findings, then inspect the Auto-Fix diff and download the report. The browser test also verified registration, explicit sign-out and sign-in, downloads, and the agent monitor. No API key was entered, and the sample workflow used the built-in deterministic engine.")

# 7. measured results
s = base("Measured results on bundled samples", "Evidence, not a broad accuracy claim", 7)
table = s.shapes.add_table(4, 5, Inches(.72), Inches(1.9), Inches(11.9), Inches(2.7)).table
widths = [2.9, 1.6, 2.7, 2.7, 2.0]
for i, w in enumerate(widths): table.columns[i].width = Inches(w)
rows = [
    ["Sample", "Cases", "Before pass/fail/error", "After pass/fail/error", "Fixes"],
    ["sample_buggy.py", "84", "77 / 6 / 1", "82 / 1 / 1", "3 verified"],
    ["sample_inventory.py", "135", "123 / 9 / 3", "131 / 1 / 3", "7 verified"],
    ["sample_shop.py", "26", "23 / 3 / 0", "26 / 0 / 0", "2 verified"],
]
for r, vals in enumerate(rows):
    for c, val in enumerate(vals):
        cell = table.cell(r, c); cell.text = val; cell.margin_left = Inches(.1); cell.margin_right = Inches(.06)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if r == 0 else (PALE if r % 2 else WHITE)
        for p in cell.text_frame.paragraphs:
            p.font.name = "Aptos"; p.font.size = Pt(13 if r else 12); p.font.bold = r == 0
            p.font.color.rgb = WHITE if r == 0 else INK
text(s, .95, 4.95, 11, .42, "Total: 245 cases · passing 223 → 239 · 12 accepted fixes · 0 regressions", 18, TEAL, True)
text(s, .95, 5.65, 11.3, .7, "Local trusted execution · deterministic engine · three planted-bug samples · 2 failures + 4 errors remained", 13, MUTED)
text(s, .95, 6.28, 11.3, .45, "Future scope: larger corpus · Docker evaluation · CI feedback · optional MCP / LangGraph exploration", 12, INK)
notes(s, "These are measured results from real Orchestrator.start runs, using the supplied requirements files and local trusted execution on bundled samples. Across 245 generated cases, passing outcomes increased from 223 to 239 after 12 accepted fixes. The fix agent reported zero regressions. Two failed assertions and four errors remained, so this is not a claim that all defects were fixed. The dataset contains only three intentionally buggy samples; it is too small to support a general accuracy claim. Future scope includes a larger independent corpus, Docker evaluation, CI integration, MCP-backed tools, more languages, and an optional LangGraph comparison. The exact settings, findings, timings, and limitations are documented in the validation report and JSON artifact.")

# 8. conclusion
s = base("Evidence-led testing, with clear limits", "Conclusion and next steps", 8)
rect(s, .85, 1.9, 7.2, 4.4, WHITE)
text(s, 1.22, 2.23, 6.4, .45, "What the system demonstrates", 20, NAVY, True)
text(s, 1.22, 2.95, 6.35, 2.25, "• Coordinated, bounded agent workflow\n• Confirmed and heuristic evidence labeled\n• Fixes re-tested before acceptance\n• API, frontend, and browser journey verified", 18, INK)
rect(s, 8.45, 1.9, 3.95, 4.4, PALE)
text(s, 8.8, 2.23, 3.2, .45, "Next", 20, NAVY, True)
text(s, 8.8, 2.95, 3.1, 2.2, "Broaden sample corpus\nAdd isolated execution coverage\nMeasure across independent modules", 17, INK)
notes(s, "TestPilot AI demonstrates a complete evidence path from source to report, with a bounded workflow, explicit oracle quality, and patch verification. The browser and endpoint checks cover the client-to-service integration, while the three sample runs show actual before and after outcomes. The main next step is evaluation on a larger, independently selected corpus and validation under the recommended Docker isolation mode. We should continue to report unresolved cases and avoid treating heuristic success as proof. Thank you; I am ready for questions.")

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)

# Re-open the generated artifact, verify slide/notes structure, and print titles.
check = Presentation(OUT)
assert 1 <= len(check.slides) <= 8, f"Unexpected slide count: {len(check.slides)}"
expected_titles = ["TestPilot AI", "The testing gap", "From single app to contract-driven system",
                   "Seven agents, one bounded workflow", "System architecture", "Live flow: sample to verified patch",
                   "Measured results on bundled samples", "Evidence-led testing, with clear limits"]
titles = []
for i, slide in enumerate(check.slides, 1):
    slide_text = [shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
    actual_title = next((value for value in slide_text if value in expected_titles), None)
    assert actual_title is not None, f"Slide {i} title was not found in the reopened deck: {slide_text}"
    titles.append(actual_title)
    note = slide.notes_slide.notes_text_frame.text.strip()
    assert len(note) >= 100, f"Slide {i} has missing/short speaker notes"
assert titles == expected_titles, f"Unexpected presentation title sequence: {titles}"
print(f"python-pptx validation: {len(check.slides)} slides; notes present on every slide")
for i, title in enumerate(titles, 1):
    print(f"{i}. {title.encode('ascii', 'replace').decode('ascii')}")
print(f"Saved: {OUT}")
